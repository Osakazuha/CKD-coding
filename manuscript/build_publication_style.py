from __future__ import annotations

import re
from pathlib import Path
import sys

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from manuscript.paths import SOURCE  # noqa: E402

FONT = "Times New Roman"
INK = RGBColor(20, 20, 20)
NAVY = RGBColor(23, 54, 93)
MUTED = RGBColor(85, 85, 85)
RULE = "808080"
LIGHT_RULE = "C8C8C8"
PAGE_DXA = 10224
COL_DXA = 4932


def _set_font(run, size=9.5, *, bold=None, italic=None, color=INK):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    fonts.set(qn("w:eastAsia"), FONT)
    fonts.set(qn("w:cs"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _set_style_font(style, size, *, bold=False, color=INK):
    style.font.name = FONT
    rpr = style._element.get_or_add_rPr()
    fonts = rpr.rFonts
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    fonts.set(qn("w:eastAsia"), FONT)
    fonts.set(qn("w:cs"), FONT)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color


def _add_field(paragraph, instruction):
    run = paragraph.add_run()
    for kind, value in (("begin", None), (None, instruction), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
        elif value == instruction:
            node = OxmlElement("w:instrText")
            node.set(qn("xml:space"), "preserve")
            node.text = value
        else:
            node = OxmlElement("w:t")
            node.text = value
        run._r.append(node)
    _set_font(run, size=7.5, color=MUTED)


def _set_columns(section, count=1, gap=360):
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), str(gap if count > 1 else 0))
    cols.set(qn("w:equalWidth"), "1")


def _remove_doc_grid(section):
    grid = section._sectPr.find(qn("w:docGrid"))
    if grid is not None:
        section._sectPr.remove(grid)


def _configure_page(section, *, columns=1):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.70)
    section.bottom_margin = Inches(0.70)
    section.left_margin = Inches(0.70)
    section.right_margin = Inches(0.70)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.35)
    _set_columns(section, columns)
    _remove_doc_grid(section)


def _configure_header_footer(section, *, short_title="", link=False):
    section.header.is_linked_to_previous = link
    section.footer.is_linked_to_previous = link
    if not link:
        hp = section.header.paragraphs[0]
        hp.text = short_title
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hp.paragraph_format.space_after = Pt(0)
        for run in hp.runs:
            _set_font(run, size=7.5, italic=True, color=MUTED)
        fp = section.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.space_after = Pt(0)
        _add_field(fp, "PAGE")


def _configure_styles(doc):
    normal = doc.styles["Normal"]
    _set_style_font(normal, 9.35)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.06
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(1.8)
    normal.paragraph_format.widow_control = True

    for name, size, before, after, color in (
        ("Heading 1", 11.5, 8, 3, NAVY),
        ("Heading 2", 10.0, 6, 2, INK),
        ("Heading 3", 9.5, 4, 1.5, INK),
    ):
        style = doc.styles[name]
        _set_style_font(style, size, bold=True, color=color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for name in ("Figure Caption", "Table Caption"):
        if name not in doc.styles:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = doc.styles[name]
        _set_style_font(style, 7.2, color=INK)
        style.font.italic = name == "Figure Caption"
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.space_before = Pt(2)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_together = True


TOKEN_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def _add_inline(paragraph, text, size=9.5):
    pos = 0
    for match in TOKEN_RE.finditer(text):
        if match.start() > pos:
            _set_font(paragraph.add_run(text[pos : match.start()]), size=size)
        token = match.group(0)
        if token.startswith("**"):
            _set_font(paragraph.add_run(token[2:-2]), size=size, bold=True)
        else:
            _set_font(paragraph.add_run(token[1:-1]), size=size, italic=True)
        pos = match.end()
    if pos < len(text):
        _set_font(paragraph.add_run(text[pos:]), size=size)


def _add_body(doc, text, *, indent=True, size=9.35, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.first_line_indent = Inches(0.16) if indent else Inches(0)
    p.paragraph_format.line_spacing = 1.06
    p.paragraph_format.space_after = Pt(1.8)
    _add_inline(p, text, size=size)
    return p


def _add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
    _add_inline(p, text, size={1: 11.5, 2: 10.0, 3: 9.5}.get(level, 9.5))
    for run in p.runs:
        run.bold = True
        if level == 1:
            run.font.color.rgb = NAVY
    return p


def _parse_table(lines, start):
    rows = []
    idx = start
    while idx < len(lines) and lines[idx].lstrip().startswith("|"):
        cells = [cell.strip() for cell in lines[idx].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        idx += 1
    return rows, idx


def _set_cell_margins(cell, top=32, start=45, bottom=32, end=45):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_borders(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, attrs in edges.items():
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        for key, value in attrs.items():
            node.set(qn(f"w:{key}"), str(value))


def _table_widths(rows):
    ncols = len(rows[0])
    if ncols == 2:
        weights = [0.68, 0.32]
    elif ncols == 3:
        weights = [0.53, 0.235, 0.235]
    elif ncols == 4:
        weights = [0.42, 0.20, 0.25, 0.13]
    else:
        weights = [1 / ncols] * ncols
    widths = [round(COL_DXA * weight) for weight in weights]
    widths[-1] += COL_DXA - sum(widths)
    return widths


def _set_table_geometry(table, widths):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(COL_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "45")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        node = OxmlElement("w:gridCol")
        node.set(qn("w:w"), str(width))
        grid.append(node)
    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for idx, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _add_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    widths = _table_widths(rows)
    no_border = {"val": "nil"}
    top_rule = {"val": "single", "sz": "8", "color": RULE}
    header_rule = {"val": "single", "sz": "6", "color": RULE}
    bottom_rule = {"val": "single", "sz": "8", "color": RULE}
    light_rule = {"val": "single", "sz": "2", "color": LIGHT_RULE}

    for idx, value in enumerate(rows[0]):
        cell = table.rows[0].cells[idx]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        _add_inline(p, value, size=6.7)
        for run in p.runs:
            run.bold = True
        _set_cell_borders(cell, top=top_rule, bottom=header_rule, left=no_border, right=no_border)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)

    for row_idx, values in enumerate(rows[1:], 1):
        row = table.add_row()
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        cells = row.cells
        for idx, value in enumerate(values):
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = row_idx < len(rows) - 1
            _add_inline(p, value, size=6.5)
            is_last = row_idx == len(rows) - 1
            _set_cell_borders(
                cells[idx],
                bottom=bottom_rule if is_last else light_rule,
                left=no_border,
                right=no_border,
            )
    _set_table_geometry(table, widths)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    return table


def _add_image(doc, line, caption_text=None, *, source_dir=None):
    match = re.match(r"!\[(.*?)\]\((.*?)\)", line)
    if not match:
        return
    alt_text, relative = match.groups()
    path = ((source_dir or SOURCE.parent) / relative).resolve()
    if not path.is_file():
        raise FileNotFoundError(f"manuscript figure is missing: {path}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Inches(3.24))
    drawings = run._r.xpath(".//wp:docPr")
    if drawings:
        drawings[0].set("descr", caption_text or alt_text)
    if caption_text:
        cp = doc.add_paragraph(style="Figure Caption")
        cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _add_inline(cp, caption_text, size=7.2)


def _add_front_matter(doc, lines, title_end):
    title = lines[0].strip()[2:]
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.02
    _set_font(p.add_run(title), size=17.0, bold=True, color=NAVY)
    for line in lines[1:title_end]:
        text = line.strip()
        if not text:
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(2)
        _add_inline(p, text, size=9.5 if "Lening Yang" in text else 8.6)


def build(source: Path = SOURCE, output: Path | None = None) -> Path:
    """Render one Markdown manuscript source to DOCX; figures resolve beside the source."""
    source = Path(source)
    output = Path(output) if output is not None else source.with_suffix(".docx")
    lines = source.read_text(encoding="utf-8").splitlines()
    if not lines or not lines[0].startswith("# "):
        raise ValueError("Publication markdown must begin with a level-one title")
    abstract_idx = next(i for i, line in enumerate(lines) if line.strip() == "## Abstract")

    doc = Document()
    doc.core_properties.title = lines[0][2:].strip()
    doc.core_properties.author = "Lening Yang"
    doc.core_properties.subject = "Leakage-controlled CKD machine-learning evaluation"
    doc.core_properties.keywords = "chronic kidney disease; machine learning; data leakage; PCA; SHAP; calibration"
    doc.core_properties.comments = "Publication-style scientific article"
    _configure_styles(doc)
    _configure_page(doc.sections[0], columns=1)
    _configure_header_footer(doc.sections[0], short_title="")
    doc.sections[0].different_first_page_header_footer = True
    _add_front_matter(doc, lines, abstract_idx)

    current_section = ""
    body_columns_started = False
    idx = abstract_idx
    while idx < len(lines):
        line = lines[idx].strip()
        if not line:
            idx += 1
            continue

        if line.startswith("# "):
            idx += 1
            continue

        if line.startswith("## "):
            heading = line[3:].strip()
            if heading.startswith("1. Introduction") and not body_columns_started:
                section = doc.add_section(WD_SECTION.NEW_PAGE)
                _configure_page(section, columns=2)
                section.different_first_page_header_footer = False
                _configure_header_footer(section, short_title="Leakage-controlled CKD classification", link=False)
                body_columns_started = True
            current_section = heading
            _add_heading(doc, heading, 1)
            idx += 1
            continue

        if line.startswith("### "):
            heading = line[4:].strip()
            if re.match(r"(Supplementary\s+)?Table\s+S?\d+\.", heading):
                cp = doc.add_paragraph(style="Table Caption")
                cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                _add_inline(cp, heading, size=7.2)
            else:
                _add_heading(doc, heading, 2)
            idx += 1
            continue

        if line.startswith("!["):
            next_idx = idx + 1
            while next_idx < len(lines) and not lines[next_idx].strip():
                next_idx += 1
            caption = None
            if next_idx < len(lines):
                candidate = lines[next_idx].strip()
                if candidate.startswith("*Figure ") and candidate.endswith("*"):
                    caption = candidate[1:-1]
            _add_image(doc, line, caption, source_dir=source.parent)
            idx = next_idx + 1 if caption else idx + 1
            continue

        if line.startswith("|"):
            rows, idx = _parse_table(lines, idx)
            if rows:
                _add_table(doc, rows)
            continue

        if line.startswith("**Keywords:"):
            p = _add_body(doc, line, indent=False, size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(5)
        elif current_section == "Abstract":
            p = _add_body(doc, line, indent=False, size=9.2)
            p.paragraph_format.space_after = Pt(3)
        elif current_section == "References":
            p = _add_body(doc, line, indent=False, size=7.0, align=WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.left_indent = Inches(0.13)
            p.paragraph_format.first_line_indent = Inches(-0.13)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(1.4)
        elif current_section == "Acknowledgments":
            _add_body(doc, line, indent=False, size=8.8)
        elif line.startswith("*") and line.endswith("*"):
            p = _add_body(doc, line, indent=False, size=7.2, align=WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(4)
        else:
            _add_body(doc, line, indent=True)
        idx += 1

    for section in doc.sections:
        _remove_doc_grid(section)
        line_number = section._sectPr.find(qn("w:lnNumType"))
        if line_number is not None:
            section._sectPr.remove(line_number)

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--output", type=Path, required=True, help="DOCX path (a run-local publication stage; never a project-level path)")
    arguments = parser.parse_args()
    print(build(arguments.source, arguments.output))
